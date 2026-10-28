"""Markdown 表格解析與轉出（HTML / DOCX）。"""

from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt


@dataclass
class MdTable:
    headers: list[str]
    rows: list[list[str]]
    aligns: list[str]  # left / center / right


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def _parse_aligns(sep_cells: list[str]) -> list[str]:
    aligns: list[str] = []
    for cell in sep_cells:
        c = cell.strip().replace(" ", "")
        if c.startswith(":") and c.endswith(":"):
            aligns.append("center")
        elif c.endswith(":"):
            aligns.append("right")
        else:
            aligns.append("left")
    return aligns


def parse_markdown_tables(text: str) -> list[MdTable]:
    """從 Markdown 文字中找出所有 GFM 表格。"""
    tables: list[MdTable] = []
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    while i < len(lines) - 1:
        header_line = lines[i].strip()
        sep_line = lines[i + 1].strip() if i + 1 < len(lines) else ""
        if (
            header_line.startswith("|")
            and "|" in header_line[1:]
            and re.match(r"^\|?[\s\-:|]+\|[\s\-:|]*\|?$", sep_line)
        ):
            headers = _split_row(header_line)
            aligns = _parse_aligns(_split_row(sep_line))
            # 對齊數量對齊欄位數
            while len(aligns) < len(headers):
                aligns.append("left")
            aligns = aligns[: len(headers)]

            rows: list[list[str]] = []
            j = i + 2
            while j < len(lines):
                row_line = lines[j].strip()
                if not row_line.startswith("|") or "|" not in row_line[1:]:
                    break
                cells = _split_row(row_line)
                # 補齊 / 截斷欄位
                if len(cells) < len(headers):
                    cells.extend([""] * (len(headers) - len(cells)))
                cells = cells[: len(headers)]
                rows.append(cells)
                j += 1

            tables.append(MdTable(headers=headers, rows=rows, aligns=aligns))
            i = j
            continue
        i += 1
    return tables


def table_to_html(table: MdTable) -> str:
    """轉成可貼到 Word 的 HTML table。"""
    parts = [
        '<table border="1" cellspacing="0" cellpadding="4" '
        'style="border-collapse:collapse;border:1px solid #000;">'
    ]
    parts.append("<thead><tr>")
    for idx, h in enumerate(table.headers):
        align = table.aligns[idx] if idx < len(table.aligns) else "left"
        parts.append(
            f'<th style="border:1px solid #000;padding:4px 8px;'
            f'background:#f0f0f0;text-align:{align};">{html.escape(h)}</th>'
        )
    parts.append("</tr></thead><tbody>")
    for row in table.rows:
        parts.append("<tr>")
        for idx, cell in enumerate(row):
            align = table.aligns[idx] if idx < len(table.aligns) else "left"
            parts.append(
                f'<td style="border:1px solid #000;padding:4px 8px;'
                f'text-align:{align};">{html.escape(cell)}</td>'
            )
        parts.append("</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


def tables_to_html(tables: list[MdTable]) -> str:
    if not tables:
        return ""
    return "<br/><br/>".join(table_to_html(t) for t in tables)


def table_to_plain(table: MdTable) -> str:
    lines = ["\t".join(table.headers)]
    for row in table.rows:
        lines.append("\t".join(row))
    return "\n".join(lines)


def tables_to_plain(tables: list[MdTable]) -> str:
    return "\n\n".join(table_to_plain(t) for t in tables)


def _set_cell_border(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def save_tables_to_docx(tables: list[MdTable], path: str | Path) -> Path:
    """將表格寫入 .docx 檔。"""
    path = Path(path)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft JhengHei"
    style.font.size = Pt(11)

    for t_idx, table in enumerate(tables):
        if t_idx > 0:
            doc.add_paragraph("")

        cols = len(table.headers)
        doc_table = doc.add_table(rows=1 + len(table.rows), cols=cols)
        doc_table.style = "Table Grid"

        for c, header in enumerate(table.headers):
            cell = doc_table.rows[0].cells[c]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
            _set_cell_border(cell)

        for r, row in enumerate(table.rows):
            for c, value in enumerate(row):
                cell = doc_table.rows[r + 1].cells[c]
                cell.text = value
                _set_cell_border(cell)

    doc.save(path)
    return path
